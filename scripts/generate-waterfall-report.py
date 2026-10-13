"""Laporan baru dari format DOCX peserta, source proyek dan bukti uji aktual."""
from pathlib import Path
from copy import deepcopy
from collections import defaultdict
from datetime import datetime, timezone, timedelta
import json
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "docs/laporan-waterfall"
OUTPUT = ROOT / "pengumpulan/07-Laporan-Proyek-Waterfall.docx"
TEMPLATE = DATA / "FORMAT-LAPORAN.docx"
NAME, NPM = "M. Sidiq Firdaus", "2315061118"


def prune_images(document):
    used = {node.get(attribute) for node in document.element.iter() for attribute in (qn("r:embed"), qn("r:id")) if node.get(attribute)}
    for rid, rel in list(document.part.rels.items()):
        if rel.reltype.endswith(("/image", "/hyperlink")) and rid not in used:
            document.part.drop_rel(rid)


def prepare_template():
    if TEMPLATE.exists():
        return
    original = ROOT / "laporan-proyek-sistem-restoran.docx"
    document = Document(original)
    elements = list(document.element.body)
    first_chapter = next(i for i, element in enumerate(elements) if "BAB I" in "".join(element.itertext()))
    for element in elements[first_chapter:]:
        if element.tag != qn("w:sectPr"):
            document.element.body.remove(element)
    replacements = {
        "BAGAS PANGESTU": NAME.upper(), "NPM 2315061010": f"NPM {NPM}",
        "RANCANG BANGUN SISTEM RESERVASI & PEMESANAN RESTORAN": "RANCANG BANGUN SISTEM RESERVASI & PEMESANAN RESTORAN\nRASA & RUANG",
        "JURUSAN TEKNIK ELEKTRO": "", "FAKULTAS TEKNIK": "",
    }
    for paragraph in document.paragraphs:
        if paragraph.text in replacements:
            value = replacements[paragraph.text]
            if paragraph.runs:
                paragraph.runs[0].text = value
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run(value)
    document.core_properties.author = NAME
    document.core_properties.last_modified_by = NAME
    document.core_properties.title = "Format Laporan Rasa & Ruang"
    prune_images(document)
    document.save(TEMPLATE)


# Hilangkan singkatan rapat pada narasi/tabel, tanpa mengubah code block.
REPLACEMENTS = {
    "npmci": "npm ci", "npmrunseed:demo": "npm run seed:demo", "npmrun debug": "npm run debug", "npmrun": "npm run",
    "Node24": "Node 24", "ENV_USER": "ENV_USER", "clocktetap": "clock tetap", "DBsementara": "DB sementara",
    "DBapp": "DB aplikasi", "DBaplikasi": "DB aplikasi", "dataapp": "data aplikasi", "Server/DBaktif": "Server/DB aktif",
    "menuaktif": "menu aktif", "tahunUTC": "tahun UTC", "inputkalender": "input kalender", "dailykosong": "daily kosong",
    "firstmenuID": "ID menu pertama", "schemaSQL": "schema SQL", "snapshotlama": "snapshot lama", "reportCSV": "report CSV",
    "hargaDB": "harga DB", "harga35000": "harga 35000", "table_id3": "table_id 3", "total105000": "total 105000",
    "statusawal": "status awal", "read-only": "read-only", "di awal": "di awal", "Source": "Source",
    "confirmedonlyupdate": "hanya confirmed boleh diubah", "confirmed/seated": "confirmed / seated",
    "SQLberparameter": "SQL berparameter", "statusillegal": "status ilegal", "statusbuttons": "tombol status",
    "duplicatephone": "telepon duplikat", "teleponunik": "telepon unik", "UNIQUEphone": "UNIQUE phone",
    "bodykosong": "body kosong", "bodybesar": "body besar", "malformedJSON": "malformed JSON", "JSONparser": "JSON parser",
    "date.min=today": "date.min = today", "withclock": "dengan clock", "clock hari": "clock hari",
    "scenariopositif": "skenario positif", "nilaiuang": "nilai uang", "nilai/status": "nilai / status",
    "summaryakhir": "summary akhir", "summarynol": "summary nol", "buktitersedia": "bukti tersedia",
    "namedresponses": "named responses", "namedresponse": "named response", "unit/service": "unit/service",
    "callback selesai": "callback selesai", "writes.length0": "writes.length = 0", "logbegin": "log begin",
    "clock2026": "clock 2026", "hari sama00": "hari sama 00", "customer dibuatdi awal": "customer dibuat di awal",
    "berspasi": "berspasi", "Date(start_at)": "Date(start_at)", "tamumax": "tamu maksimum",
    "tidakada": "tidak ada", "tidakmenampilkan": "tidak menampilkan", "tidakmengakui": "tidak mengakui",
    "tidak bocor": "tidak bocor", "rollBack": "rollBack", "regresiterkait": "regresi terkait",
    "mendatangWIB": "mendatang WIB", "kalenderinvalid": "kalender invalid", "start==now": "start == now",
    "transisiillegal": "transisi ilegal", "inputinvalid": "input invalid", "errorstring": "error string",
    "header/harian/TOTAL": "header / harian / TOTAL", "lolos validasi bentuk": "lolos validasi bentuk",
    "MAXmenutotal": "maksimum menu", "tanggalkosong": "tanggal kosong", "tanggalbesokUTCpukul": "tanggal besok UTC pukul ",
    "hashsnapshot": "hash snapshot", "tanpaformat": "tanpa format", "DefaultCSV": "Default CSV", "JSONsubclass": "subclass JSON",
    "readonly": "read-only", "Tersedia/Dipesan/Tamu hadir": "Tersedia / Dipesan / Tamu hadir",
    "submitlaporan": "submit laporan", "constform": "const form", "akhir2100": "akhir 2100", "awal2000": "awal 2000",
    "archiver7/globdeprecated": "archiver 7 / glob deprecated", "archiver8+ZipArchive": "archiver 8 + ZipArchive",
    "auditregistrybaru": "audit registry baru", "ZIP allowlist": "ZIP allowlist", "memiliki riwayat": "memiliki riwayat",
    "Clocktetap": "Clock tetap", "V: ": "V: ", "fotoform": "foto form", "operatorlokal": "operator lokal",
    "PascalCaseclass": "PascalCase class", "caller": "caller", "dependencySQL": "dependency SQL",
    "payloadarray": "payload array", "kapasitaskurang": "kapasitas kurang", "Notes501": "Notes 501", "DBuji": "DB uji",
    "serveruji": "server uji", "Serveruji": "Server uji", "stubdriver": "stub driver", "Datauji": "Data uji",
    "unit.tap/integration.tap": "unit.tap / integration.tap", "TAP/JSON": "TAP / JSON",
    "firstmenu": "menu pertama", "HTTP45": "HTTP 45", "HTTP21": "HTTP 21",
    "bilanganbulat": "bilangan bulat", "campur": "campur", "from/to": "from / to", "from≤to": "from ≤ to",
    "ChangeStatus": "ChangeStatus", "URIrequest": "URI request", "Syntaxcheck": "Syntax check", "syntaxcheck": "syntax check",
    "hanyalog": "hanya log", "tidak respons": "tidak respons", "aktifDB": "aktif DB",
    "poolditutup": "pool ditutup", "loggagal": "log gagal", "write rollback": "write rollback",
    "slotlama": "slot lama", "slotbaru": "slot baru", "table1": "table 1", "menu1": "menu 1", "phoneberhuruf": "phone berhuruf",
    "Factory": "Factory", "baru13": "baru 13", "PHPUnit": "PHPUnit", "syntax24": "syntax 24",
    "lihatfilter": "lihat / filter", "parse2000": "parse 2000", "fixtureterisolasi": "fixture terisolasi",
    "DBport": "DB port", "servicedihentikan": "service dihentikan", "poolend": "pool.end", "sama-sama": "sama-sama",
    "errorrollback": "error rollback", "periode11": "periode 11", "runtimefrontend": "runtime frontend",
    "sudahada": "sudah ada", "primarysource": "primary source", "stateUI": "state UI",
    "kapasitasguests": "kapasitas guests", "portsementara": "port sementara", "servicefake": "service fake",
    "IDsaved": "ID saved", "DBseedterisolasi": "DB seed terisolasi", "observasiUI": "observasi UI",
    "CSVbrowser": "CSV browser", "mobilebaru": "mobile baru", "XSSend-to-end": "XSS end-to-end",
    "rehearsalglobal": "rehearsal global", "belumada": "belum ada", "salin.env.example": "salin .env.example",
    "periksaDB_HOST": "periksa DB_HOST", "sertaHOST": "serta HOST", "TanggalbesokUTCpukul": "Tanggal besok UTC pukul ",
    "aplikasibentrok": "aplikasi bentrok", "customerexisting": "customer existing", "kontaklama": "kontak lama",
    "POSTphoneexisting": "POST phone existing", "IDcustomer": "ID customer", "IDmenu": "ID menu",
    "menuinactive": "menu inactive", "sudahBEGIN": "sudah BEGIN", "ulangreferensi": "ulang referensi",
    "yangada": "yang ada", "InactiveDBbaru": "Inactive DB baru", "servicecapacity": "service capacity",
    "doublebook": "double booking", "Requestadjacent": "Request adjacent", "mejalebihbesar": "meja lebih besar",
    "IDtidakditemukan": "ID tidak ditemukan", "IDtidakada": "ID tidak ada", "setelahdiproses": "setelah diproses",
    "Janganubah": "Jangan ubah", "terminaldibuka": "terminal dibuka", "statusunknown": "status unknown",
    "modelchangeStatus": "model changeStatus", "DELETEcustomer": "DELETE customer", "punyareservasi": "punya reservasi",
    "memaksaDELETE": "memaksa DELETE", "serviceaffectedRows": "service affectedRows", "querytanggalhilang": "query tanggal hilang",
    "tanggalvalid": "tanggal valid", "PDFmelaluiCetakbrowser": "PDF melalui Cetak browser", "GETtidak-ada": "GET tidak-ada",
    "DBtidakaktif": "DB tidak aktif", "jalankandb:setup": "jalankan db:setup", "padaDByangbenar": "pada DB yang benar",
    "janganresetdata": "jangan reset data", "servermiliksendiri": "server milik sendiri", "ubahPORT": "ubah PORT",
    "takdikenal": "tak dikenal", "pesanumum": "pesan umum", "detaillog": "detail log", "tidakmembocorkan": "tidak membocorkan",
    "Skenariostub": "Skenario stub", "masihrencana": "masih rencana", "skenariostub": "skenario stub",
    "UPDATEconflict": "UPDATE conflict", "INSERT/commiterror": "INSERT / commit error", "ulangrequest": "ulang request",
    "Driverwrite": "Driver write", "Batasaktual": "Batas aktual", "keduaerror": "kedua error", "sebelumpatch": "sebelum patch",
    "tidakterjangkau": "tidak terjangkau", "dariUI": "dari UI", "Fetchgagal": "Fetch gagal", "tombolbusy": "tombol busy",
    "Startserver": "Start server", "periksaURL": "periksa URL", "setelahwrite": "setelah write", "menampilkanerror": "menampilkan error",
    "tidakmenganggap": "tidak menganggap", "refresherror": "refresh error", "tampilkanlaporan": "tampilkan laporan",
    "Tampilkanlaporan": "Tampilkan laporan", "sebelumawait": "sebelum await", "setelahawait": "setelah await",
    "Gunakanreferensi": "Gunakan referensi", "submitreport": "submit report", "Bugnyata": "Bug nyata", "telahfix": "telah diperbaiki",
    "observasiUIbaru": "observasi UI baru", "meja4kursi": "meja 4 kursi", "maxintegerlolos": "max integer lolos",
}


def pretty(text):
    # Jangan menyisipkan spasi pada URL, ISO datetime, nama library atau tombol debugger.
    protected = []
    def protect(match):
        protected.append(match[0])
        return "\ufff0" + str(len(protected) - 1) + "\ufff1"
    text = re.sub(r"https?://\S+|\d{4}-\d{2}-\d{2}T\d{2}:\d{2}(?::\d{2}(?:\.\d+)?Z?)?|\b(?:mysql2|3NF|F5|F10|F11)\b", protect, text)
    for source, target in REPLACEMENTS.items():
        text = text.replace(source, target)
    text = re.sub(r"(?<=\d)(?=[A-Za-z])|(?<=[A-Za-z])(?=\d)", " ", text)
    for i, value in enumerate(protected):
        text = text.replace("\ufff0" + str(i) + "\ufff1", value)
    return text


errors = [
    ("ER-01", "JSON rusak", "POST body {invalid", "400 error parser", "Perbaiki JSON/Content-Type; ApiClient menampilkan pesan. Parser Express + middleware.", "IT-022;HTTP-034"),
    ("ER-02", "Body melebihi32KB", "POST data33000karakter", "413; tidak diproses service", "Kurangi payload; express.json(limit32kb).", "IT-022"),
    ("ER-03", "Payload bukan objek", "null/[]", "400", "Kirim objectJSON; requireObject.", "UT-022/023;IT-004;HTTP-033"),
    ("ER-04", "Nama kosong/terlalu panjang", "Whitespace atau>100karakter", "400", "Isi/trim nama sesuai batas; text/validateCustomer.", "UT-018;HTTP-032;BT-001"),
    ("ER-05", "Telepon invalid", "Huruf, <8/>15digit setelah normalisasi", "400", "Gunakan digit dengan +opsional; inputawal≤20karakter; validateCustomer.", "UT-019/020;BT-002"),
    ("ER-06", "Email invalid", "dina@ atau>150karakter", "400", "Kosongkan jika opsional atau perbaiki format/panjang.", "UT-021;BT-003"),
    ("ER-07", "Telepon duplikat", "POSTphoneexisting", "409; kontaklama tetap", "Gunakan customerexisting/update; UNIQUEphone dan ER_DUP_ENTRY.", "IT-003;HTTP-008"),
    ("ER-08", "Tanggal/jam kalender invalid", "2027-02-30, jam25, format salah/tahun di luar2000–2100", "400", "Perbaiki formatWIB/tanggal; parseDateTime round-trip.", "UT-027–029/043/062–063;HTTP-036"),
    ("ER-09", "Start lampau/sama now", "start<=clock", "400", "Pilih waktu mendatang; validateReservation.", "UT-026/064"),
    ("ER-10", "Angka/batas invalid", "guestsstring/pecahan/0/9;duration29/241;IDinvalid", "400", "Gunakan numberinteger dan rentang; integer helper.", "UT-030–035;HTTP-035/042;BT-004/005/009"),
    ("ER-11", "Pesanan invalid", "itemsbukanarray,duplikat,qty0/100,>30menu,notes>500", "400", "Hapus duplikat, sesuaikanqty/array/catatan; validateReservation.", "UT-036–040;HTTP-037;BT-006–008"),
    ("ER-12", "Referensi tidak ada/aktif", "IDcustomer/meja/menu tidak ada atau menuinactive", "404; rollback jika sudahBEGIN", "Muat ulangreferensi/pilih data yangada; service lookup/FK. InactiveDBbaru rencana.", "UT-002/004/006;IT-009;HTTP-013;BT-010"),
    ("ER-13", "Kapasitas meja kurang", "guests5 di meja4kursi", "400; no write/rollback", "Pilih meja lebihbesar atau kurangi tamu; servicecapacity.", "UT-003;IT-007;HTTP-011"),
    ("ER-14", "Jadwal overlap/concurrent", "Meja/intervalaktif sama", "409; tidak doublebook", "Pilih waktu/meja lain; lock+intervalquery. Requestadjacent boleh.", "UT-005;IT-008/010/011/025;HTTP-012"),
    ("ER-15", "Update/status reservasi missing", "IDtidakditemukan", "404", "Refresh daftar/periksa ID; service SELECT.", "UT-009/013;HTTP-043"),
    ("ER-16", "Update selain confirmed", "PUT seated/completed/cancelled", "409", "Janganubah detail setelahdiproses; ikuti alur status; serviceexisting.", "UT-010;IT-016;HTTP-019"),
    ("ER-17", "Transisi status ilegal", "confirmed→completed; terminaldibuka;statusunknown", "409; state tetap", "Ikuti confirmed→seated→completed atau cancel; modelchangeStatus.", "UT-012/054–056/059;IT-014;HTTP-015/016/025"),
    ("ER-18", "Hapus customer ber-riwayat", "DELETEcustomer punyareservasi", "409; FK menjaga data", "Pertahankan pelanggan/riwayat; tidak menghapus transaksi untuk memaksaDELETE.", "IT-024;HTTP-028"),
    ("ER-19", "Customer missing", "PUT/DELETEIDtidakada", "404", "Refresh daftar; serviceaffectedRows/returnnull.", "UT-015/016;HTTP-031"),
    ("ER-20", "Rentang laporan invalid", "from>to atau querytanggalhilang", "400", "Isi tanggalvalid/urut; validateDateRange.", "UT-044/046;IT-028;HTTP-038/039"),
    ("ER-21", "Format exporter unsupported", "format=pdf", "400", "Pakai csv/json; PDFmelaluiCetakbrowser; factory.", "UT-067;IT-020;HTTP-040"),
    ("ER-22", "Endpoint/IDURL invalid", "GETtidak-ada; ID=abc", "404endpoint;400ID", "Periksa route/metode/ID; routefallback/integer.", "IT-028;HTTP-041/042"),
    ("ER-23", "DB belum siap saat startup", "DBtidakaktif/credential/schema salah", "Log;poolend;exitCode1", "Aktifkan MySQL, periksa.env, jalankandb:setup padaDByangbenar; janganresetdata. server.js.", "BT-016 belum dieksekusi"),
    ("ER-24", "Port server terpakai", "EADDRINUSE", "Loggagal;poolend;exitCode1", "Hentikan servermiliksendiri atau ubahPORT; server.on(error).", "BT-017 belum dieksekusi"),
    ("ER-25", "Driver/query internal gagal", "execute melemparerror takdikenal", "500pesanumum;detaillog", "Periksa console/koneksi/config; API tidakmembocorkan SQL. Skenariostub500 masihrencana.", "BT-018 belum dieksekusi"),
    ("ER-26", "Operasi transaksi gagal", "Menu missing/UPDATEconflict; INSERT/commiterror", "Rollback + finallyrelease", "Perbaiki sebab, cek persistensi, ulangrequest. Driverwrite/commit khusus masihrencana.", "UT-058;IT-009/013;BT-013/014"),
    ("ER-27", "Rollback juga gagal", "Operasi dan rollback melempar", "Release dicoba;errorawal dapat terganti", "Batasaktual: logerrorbaru; rekomendasi simpan keduaerror dan uji recovery sebelumpatch.", "BT-015 belum dieksekusi"),
    ("ER-28", "Server tidakterjangkau dariUI", "Fetchgagal/network", "Pesan koneksi; tombolbusy dipulihkan", "Startserver/periksaURL; ApiClient/catch/finally.", "UI-022 belum dieksekusi"),
    ("ER-29", "Refresh gagal setelahwrite sukses", "Write200/201 laluGETerror", "UI menampilkanerror; write mungkin tersimpan", "Periksa GET/DB dulu sebelum ulang create; tidakmenganggap refresherror=rollback.", "BT-019 belum dieksekusi"),
    ("ER-30", "Ekspor/cetak sebelum report tersedia", "state.report belumada", "Toast: tampilkanlaporan dahulu", "Tampilkanlaporan lalu ulang; public/app.js guard.", "UI-018 belum dieksekusi"),
    ("ER-31", "Bug async currentTarget", "event.currentTarget setelahawait=null", "Perbaikan menyimpan constform sebelumawait", "Gunakanreferensi form di finally; regresi submitreport. Bugnyata historis, source telahfix.", "BUKTI-BROWSER;observasiUIbaru"),
]


def markdown_table(rows):
    return "\n".join("| " + " | ".join(str(c).replace("|", "/").replace("\n", " / ") for c in row) + " |" for row in rows[:1]) + "\n| " + " | ".join("---" for _ in rows[0]) + " |\n" + "\n".join("| " + " | ".join(str(c).replace("|", "/").replace("\n", " / ") for c in row) + " |" for row in rows[1:])


prepare_template()
doc = Document(TEMPLATE)
doc.core_properties.author = NAME
doc.core_properties.last_modified_by = NAME
doc.core_properties.title = "Rancang Bangun Sistem Reservasi & Pemesanan Restoran Rasa & Ruang — Waterfall"
doc.core_properties.subject = "Laporan sertifikasi, error/recovery dan test scenario lengkap"
doc.core_properties.comments = "Isi berdasarkan source dan bukti terbaru; format mengikuti laporan contoh."
scenarios = json.loads((DATA / "TEST-SCENARIOS.json").read_text(encoding="utf-8"))
evidence = json.loads((DATA / "bukti/summary.json").read_text(encoding="utf-8"))
time_wib = datetime.fromisoformat(evidence["executedAt"].replace("Z", "+00:00")).astimezone(timezone(timedelta(hours=7)))
summary_rows = [["Pemeriksaan", "Hasil aktual", "Bukti"], ["Syntax", "24 berkas lulus", "syntax.txt"], ["Unit/service", "69 lulus;0 gagal/skipped", "unit.tap"], ["Integrasi", "28 subtest lulus;29 runner termasuk induk;0 gagal", "integration.tap"], ["HTTP", "45/45 status sesuai;summary akhir benar", "http.json/http.txt"], ["Unit coverage", "83,92% baris;98,35% cabang;84,31% fungsi", "unit.tap"], ["Integrasi coverage", "95,15% baris;84,43% cabang;93,67% fungsi", "integration.tap"], ["Observasi UI", "4 pemeriksaan read-only lulus;pageErrors=[]", "browser.json/screenshot"], ["Waktu verifikasi", time_wib.strftime("%d-%m-%Y %H:%M:%S WIB"), "summary.json"]]
error_rows = [["ID / pemicu", "Respons / sebab", "Penanganan / pemulihan", "Bukti / status"]]
for id_, name, trigger, response, recovery, proof in errors:
    error_rows.append([id_ + ": " + name + "\nInput: " + trigger, response, recovery, proof])
groups = defaultdict(list)
for case in scenarios:
    groups[case["group"]].append(case)
matrix_blocks = []
for number, (group, cases) in enumerate(groups.items(), 1):
    matrix_blocks += [f"### 2.{number} {group} — {len(cases)} skenario", "", f"Tabel L2.{number} Skenario {group}", ""]
    rows = [["ID / skenario", "Prasyarat / input / langkah", "Expected", "Actual / bukti"]]
    for case in cases:
        rows.append([case["id"] + "\n" + case["scenario"], "Prasyarat: " + case["prerequisite"] + "\nInput: " + case["input"] + "\nLangkah: " + case["steps"], case["expected"], case["actual"] + "\n" + case["evidence"]])
    matrix_blocks += [markdown_table(rows), ""]
content = (DATA / "KONTEN-LAPORAN.md").read_text(encoding="utf-8")
content = content.replace("{{EVIDENCE_SUMMARY}}", markdown_table(summary_rows)).replace("{{ERROR_TABLE}}", markdown_table(error_rows)).replace("{{TEST_MATRIX}}", "\n".join(matrix_blocks))
normalized = []
in_code = False
for line in content.splitlines():
    if line.startswith("```"):
        in_code = not in_code
        normalized.append(line)
    else:
        normalized.append(line if in_code else pretty(line))
content = "\n".join(normalized) + "\n"
(DATA / "LAPORAN-WATERFALL.md").write_text(content, encoding="utf-8")


def body(text, style="Body"):
    paragraph = doc.add_paragraph(text, style=style)
    paragraph.paragraph_format.widow_control = True
    return paragraph


def add_table(rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.autofit = False
    widths = {3: [3.4, 5.2, 5.4], 4: [3.0, 4.6, 3.6, 2.8]}.get(len(rows[0]), [14 / len(rows[0])] * len(rows[0]))
    for i, value in enumerate(rows[0]):
        table.rows[0].cells[i].text = value
    for row in rows[1:]:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value.replace(" / ", " / ")
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for i, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        if i == 0:
            row_properties.append(OxmlElement("w:tblHeader"))
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
            properties = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement("w:" + edge)
                for key, value in {"val": "single", "sz": "4", "color": "808080"}.items():
                    border.set(qn("w:" + key), value)
                borders.append(border)
            properties.append(borders)
            if i == 0:
                shading = OxmlElement("w:shd"); shading.set(qn("w:fill"), "D9E2F3"); properties.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pf = paragraph.paragraph_format
                pf.first_line_indent = Cm(0)
                pf.space_before = Pt(2); pf.space_after = Pt(3); pf.line_spacing = 1.0
                pf.keep_with_next = False
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9.5 if len(rows) > 15 else 10.5)
                    run.bold = i == 0
    body("", "Normal").paragraph_format.space_after = Pt(3)


def add_figure(caption, relative):
    filename = (DATA / relative).resolve()
    if not filename.is_relative_to(ROOT / "docs"):
        raise ValueError("Gambar harus berada di docs proyek.")
    from PIL import Image
    with Image.open(filename) as image:
        ratio = image.height / image.width
    width = min(14, 19 / ratio)
    paragraph = body("", "Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(filename), width=Cm(width))
    body(caption, "CaptionGambar")


lines = content.splitlines()
code = False
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip():
        i += 1; continue
    if line.startswith("```"):
        code = not code; i += 1; continue
    if code:
        paragraph = body("", "Normal")
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(line); run.font.name = "Consolas"; run.font.size = Pt(9)
    elif line.startswith("!["):
        match = re.match(r"!\[(.+)\]\((.+)\)$", line)
        add_figure(*match.groups())
    elif line.startswith("| "):
        rows = []
        while i < len(lines) and lines[i].startswith("| "):
            if not re.match(r"^\|[\s:|\-]+\|$", lines[i]):
                rows.append([cell.strip() for cell in lines[i].split("|")[1:-1]])
            i += 1
        add_table(rows)
        continue
    elif line.startswith("### "):
        paragraph = body(line[4:], "Heading 3")
        if line.startswith("### 2."):
            paragraph.paragraph_format.page_break_before = True
    elif line.startswith("## "):
        body(line[3:], "Heading 2")
    elif line.startswith("# "):
        body(line[2:], "Heading 1")
    elif line.startswith("Tabel "):
        body(line, "CaptionTabel")
    elif line.startswith("Kode "):
        paragraph = body(line, "CaptionTabel")
    elif line.startswith("- "):
        body("• " + line[2:], "List Para")
    else:
        body(line)
    i += 1

# Hindari media lampiran proyek contoh terbawa; logo institusi sampul tetap digunakan.
prune_images(doc)
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print("Laporan baru dibuat:", OUTPUT)
print(f"6 BAB; {len(errors)} skenario error; {len(scenarios)} test scenario; {len(doc.inline_shapes)} gambar.")
